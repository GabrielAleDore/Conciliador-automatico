import io
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.schemas import ReconciliationResponse

NAVY_HEADER = "1E293B"      # Azul ardósia
TEAL_HEADER = "0F766E"      # Verde petróleo
AMBER_HEADER = "B45309"     # Âmbar
RED_HEADER = "BE123C"       # Vermelho
LIGHT_GREEN_BG = "ECFDF5"

def create_excel_report(result: ReconciliationResponse) -> bytes:
    """
    Gera relatório Excel (.xlsx) com 6 abas estruturadas,
    com foco prioritário em DATA e VALOR.
    """
    wb = Workbook()
    wb.remove(wb.active)

    font_title = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
    font_regular = Font(name="Segoe UI", size=10, color="0F172A")
    font_small = Font(name="Segoe UI", size=9, color="475569")

    thin_border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1"),
    )

    double_bottom_border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='double', color="0F172A"),
    )

    currency_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'

    # ---------------------------------------------------------
    # ABA 1: RESUMO DO FECHAMENTO
    # ---------------------------------------------------------
    ws_resumo = wb.create_sheet(title="Resumo do Fechamento")
    ws_resumo.views.sheetView[0].showGridLines = True

    ws_resumo.merge_cells("A1:E1")
    title_cell = ws_resumo["A1"]
    title_cell.value = "RELATÓRIO DE CONCILIAÇÃO BANCÁRIA RIGOROSA (DATA & VALOR)"
    title_cell.font = font_title
    title_cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_resumo.row_dimensions[1].height = 36

    d = result.dashboard
    is_ok = d.status_geral == "BALANCO_FECHADO"

    ws_resumo.merge_cells("A3:E3")
    st_cell = ws_resumo["A3"]
    st_cell.value = f"STATUS GERAL: {d.status_geral} — {d.status_mensagem}"
    st_cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF" if is_ok else "991B1B")
    st_cell.fill = PatternFill(start_color="065F46" if is_ok else "FEE2E2", fill_type="solid")
    st_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_resumo.row_dimensions[3].height = 26

    resumo_rows = [
        ("Indicador Contábil", "Sistema ERP (R$)", "Extrato Santander (R$)", "Diferença (R$)", "Status"),
        ("Saldo Inicial", float(d.saldo_inicial_sistema), float(d.saldo_inicial_banco), float(Decimal(d.saldo_inicial_sistema) - Decimal(d.saldo_inicial_banco)), "OK" if d.saldo_inicial_confere else "DIVERGENTE"),
        ("Total de Entradas", float(d.total_entradas_sistema), float(d.total_entradas_banco), float(Decimal(d.total_entradas_sistema) - Decimal(d.total_entradas_banco)), "—"),
        ("Total de Saídas", float(d.total_saidas_sistema), float(d.total_saidas_banco), float(Decimal(d.total_saidas_sistema) - Decimal(d.total_saidas_banco)), "—"),
        ("Saldo Final Declarado", float(d.saldo_final_sistema), float(d.saldo_final_banco), float(Decimal(d.saldo_final_sistema) - Decimal(d.saldo_final_banco)), "OK" if d.saldo_final_confere else "DIVERGENTE"),
    ]

    for i, row in enumerate(resumo_rows):
        row_num = 5 + i
        ws_resumo.row_dimensions[row_num].height = 22
        for col_num, val in enumerate(row, start=1):
            cell = ws_resumo.cell(row=row_num, column=col_num)
            cell.value = val
            cell.border = thin_border
            if i == 0:
                cell.font = font_header
                cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = font_regular
                if col_num in (2, 3, 4):
                    cell.number_format = currency_format
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left" if col_num == 1 else "center", vertical="center")

    # ---------------------------------------------------------
    # ABA 2: CONCILIADOS (1 A 1)
    # ---------------------------------------------------------
    ws_11 = wb.create_sheet(title="1. Conciliados (1 a 1)")
    ws_11.views.sheetView[0].showGridLines = True
    headers_11 = ["Data ERP", "Valor ERP (R$)", "Descrição ERP", "Data Banco", "Valor Banco (R$)", "Histórico Banco", "Compensação"]
    ws_11.append(headers_11)
    ws_11.row_dimensions[1].height = 26
    for c_idx in range(1, len(headers_11) + 1):
        cell = ws_11.cell(row=1, column=c_idx)
        cell.font = font_header
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for pair in result.conciliados_1_1:
        e = pair.erp_item
        b = pair.banco_item
        ws_11.append([
            e.data,
            float(Decimal(e.valor)),
            e.descricao,
            b.data,
            float(Decimal(b.valor)),
            b.descricao,
            f"D+{pair.diff_dias}" if pair.diff_dias >= 0 else f"D{pair.diff_dias}",
        ])
        r = ws_11.max_row
        ws_11.row_dimensions[r].height = 20
        for col_idx in range(1, len(headers_11) + 1):
            c = ws_11.cell(row=r, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx in (2, 5):
                c.number_format = currency_format
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 4, 7):
                c.alignment = Alignment(horizontal="center", vertical="center")

    if result.conciliados_1_1:
        tot_r = ws_11.max_row + 1
        ws_11.cell(row=tot_r, column=1, value="TOTAL 1:1").font = font_bold
        c_e = ws_11.cell(row=tot_r, column=2, value=f"=SUM(B2:B{tot_r-1})")
        c_e.number_format = currency_format
        c_e.font = font_bold
        c_e.border = double_bottom_border
        c_b = ws_11.cell(row=tot_r, column=5, value=f"=SUM(E2:E{tot_r-1})")
        c_b.number_format = currency_format
        c_b.font = font_bold
        c_b.border = double_bottom_border

    # ---------------------------------------------------------
    # ABA 3: AGRUPADOS EM LOTE (1 A N)
    # ---------------------------------------------------------
    ws_lote = wb.create_sheet(title="2. Agrupados em Lote (1 a N)")
    ws_lote.views.sheetView[0].showGridLines = True
    headers_lote = ["Origem", "Data", "Valor (R$)", "Descrição / Histórico", "Status"]
    ws_lote.append(headers_lote)
    ws_lote.row_dimensions[1].height = 26
    for c_idx in range(1, len(headers_lote) + 1):
        cell = ws_lote.cell(row=1, column=c_idx)
        cell.font = font_header
        cell.fill = PatternFill(start_color=TEAL_HEADER, end_color=TEAL_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for grp in result.conciliados_lote:
        b = grp.banco_item
        ws_lote.append(["BANCO (PAI)", b.data, float(Decimal(b.valor)), b.descricao, f"Lote com {len(grp.erp_itens)} lançamentos"])
        p_row = ws_lote.max_row
        ws_lote.row_dimensions[p_row].height = 22
        for col_idx in range(1, len(headers_lote) + 1):
            c = ws_lote.cell(row=p_row, column=col_idx)
            c.font = font_bold
            c.fill = PatternFill(start_color=LIGHT_GREEN_BG, fill_type="solid")
            c.border = thin_border
            if col_idx == 3:
                c.number_format = currency_format
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 2, 5):
                c.alignment = Alignment(horizontal="center", vertical="center")

        for e in grp.erp_itens:
            ws_lote.append(["  ↳ ERP (FILHO)", e.data, float(Decimal(e.valor)), f"    {e.descricao}", "Item do Lote"])
            c_row = ws_lote.max_row
            ws_lote.row_dimensions[c_row].height = 19
            for col_idx in range(1, len(headers_lote) + 1):
                c = ws_lote.cell(row=c_row, column=col_idx)
                c.font = font_small
                c.border = thin_border
                if col_idx == 3:
                    c.number_format = currency_format
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in (1, 2, 5):
                    c.alignment = Alignment(horizontal="center", vertical="center")

    # ---------------------------------------------------------
    # ABA 4: PENDÊNCIAS SISTEMA (ERP)
    # ---------------------------------------------------------
    ws_pend = wb.create_sheet(title="3. Pendências Sistema (ERP)")
    ws_pend.views.sheetView[0].showGridLines = True
    headers_pend = ["Data Lançamento", "Valor (R$)", "Descrição / Favorecido", "Documento", "Situação"]
    ws_pend.append(headers_pend)
    ws_pend.row_dimensions[1].height = 26
    for c_idx in range(1, len(headers_pend) + 1):
        cell = ws_pend.cell(row=1, column=c_idx)
        cell.font = font_header
        cell.fill = PatternFill(start_color=AMBER_HEADER, end_color=AMBER_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for e in result.pendencias_sistema:
        ws_pend.append([e.data, float(Decimal(e.valor)), e.descricao, e.documento or "—", "NÃO COMPENSADO NO BANCO"])
        r = ws_pend.max_row
        ws_pend.row_dimensions[r].height = 20
        for col_idx in range(1, len(headers_pend) + 1):
            c = ws_pend.cell(row=r, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx == 2:
                c.number_format = currency_format
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 4, 5):
                c.alignment = Alignment(horizontal="center", vertical="center")

    if result.pendencias_sistema:
        tot_r = ws_pend.max_row + 1
        ws_pend.cell(row=tot_r, column=1, value="TOTAL PENDÊNCIAS").font = font_bold
        c_tot = ws_pend.cell(row=tot_r, column=2, value=f"=SUM(B2:B{tot_r-1})")
        c_tot.number_format = currency_format
        c_tot.font = font_bold
        c_tot.border = double_bottom_border

    # ---------------------------------------------------------
    # ABA 5: NÃO LANÇADOS NO SISTEMA (EXCLUSIVOS BANCO)
    # ---------------------------------------------------------
    ws_banco = wb.create_sheet(title="4. Exclusivos Santander")
    ws_banco.views.sheetView[0].showGridLines = True
    headers_banco = ["Data", "Valor (R$)", "Histórico", "Categoria", "Documento"]
    ws_banco.append(headers_banco)
    ws_banco.row_dimensions[1].height = 26
    for c_idx in range(1, len(headers_banco) + 1):
        cell = ws_banco.cell(row=1, column=c_idx)
        cell.font = font_header
        cell.fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for b in result.exclusivos_banco:
        ws_banco.append([b.data, float(Decimal(b.valor)), b.descricao, b.categoria or "DÉBITO DIRETO", b.documento or "—"])
        r = ws_banco.max_row
        ws_banco.row_dimensions[r].height = 20
        for col_idx in range(1, len(headers_banco) + 1):
            c = ws_banco.cell(row=r, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx == 2:
                c.number_format = currency_format
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 4, 5):
                c.alignment = Alignment(horizontal="center", vertical="center")

    if result.exclusivos_banco:
        tot_r = ws_banco.max_row + 1
        ws_banco.cell(row=tot_r, column=1, value="TOTAL BANCO").font = font_bold
        c_tot = ws_banco.cell(row=tot_r, column=2, value=f"=SUM(B2:B{tot_r-1})")
        c_tot.number_format = currency_format
        c_tot.font = font_bold
        c_tot.border = double_bottom_border

    # ---------------------------------------------------------
    # ABA 6: DIVERGÊNCIAS DE VALORES
    # ---------------------------------------------------------
    ws_div = wb.create_sheet(title="5. Divergências de Valores")
    ws_div.views.sheetView[0].showGridLines = True
    headers_div = ["Data ERP", "Valor ERP (R$)", "Data Banco", "Valor Banco (R$)", "Diferença (R$)", "Descrição ERP", "Histórico Banco"]
    ws_div.append(headers_div)
    ws_div.row_dimensions[1].height = 26
    for c_idx in range(1, len(headers_div) + 1):
        cell = ws_div.cell(row=1, column=c_idx)
        cell.font = font_header
        cell.fill = PatternFill(start_color=RED_HEADER, end_color=RED_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for div in result.divergencias_valor:
        e = div.erp_item
        b = div.banco_item
        ws_div.append([
            e.data,
            float(Decimal(e.valor)),
            b.data,
            float(Decimal(b.valor)),
            float(Decimal(div.diff_valor)),
            e.descricao,
            b.descricao,
        ])
        r = ws_div.max_row
        ws_div.row_dimensions[r].height = 20
        for col_idx in range(1, len(headers_div) + 1):
            c = ws_div.cell(row=r, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx in (2, 4, 5):
                c.number_format = currency_format
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 3):
                c.alignment = Alignment(horizontal="center", vertical="center")

    # Ajuste automático de larguras
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and len(val_str) < 60:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
